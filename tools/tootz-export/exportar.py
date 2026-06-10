"""
Pos-processamento inspirado no repo contrattz:
- opcionalmente insere logo no cabecalho
- adiciona rodape com dados juridicos + paginacao
- linha decorativa roxa acima do rodape

Uso:
  python3 exportar.py <arquivo.docx> [logo.png]
"""

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FOOTER_TEXT = (
    "TOOTZ SOLUCOES TECNOLOGICA LTDA - ME  •  "
    "CNPJ 11.974.262/0001-73  •  "
    "atendimento@tootz.com.br"
)
PURPLE = "8B2DF7"
DARK = RGBColor(0x17, 0x16, 0x1B)
LOGO_WIDTH_CM = 3
FOOTER_FONT_PT = 8
RIGHT_TAB_TWIPS = 9070


def _set_right_tab_stop(paragraph, twips=RIGHT_TAB_TWIPS):
    p_pr = paragraph._p.get_or_add_pPr()
    for existing in p_pr.findall(qn("w:tabs")):
        p_pr.remove(existing)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(twips))
    tabs.append(tab)
    p_pr.append(tabs)


def _add_purple_top_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    for existing in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(existing)
    p_bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:color"), PURPLE)
    top.set(qn("w:space"), "4")
    p_bdr.append(top)
    p_pr.append(p_bdr)


def _add_field(paragraph, field_code, font_size_pt=FOOTER_FONT_PT):
    p = paragraph._p

    def make_rpr():
        r_pr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(font_size_pt * 2))
        r_pr.append(sz)
        return r_pr

    r_begin = OxmlElement("w:r")
    r_begin.append(make_rpr())
    fc_begin = OxmlElement("w:fldChar")
    fc_begin.set(qn("w:fldCharType"), "begin")
    r_begin.append(fc_begin)
    p.append(r_begin)

    r_instr = OxmlElement("w:r")
    r_instr.append(make_rpr())
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    r_instr.append(instr)
    p.append(r_instr)

    r_end = OxmlElement("w:r")
    r_end.append(make_rpr())
    fc_end = OxmlElement("w:fldChar")
    fc_end.set(qn("w:fldCharType"), "end")
    r_end.append(fc_end)
    p.append(r_end)


def _styled_run(paragraph, text, font_size_pt=FOOTER_FONT_PT, color=DARK):
    run = paragraph.add_run(text)
    run.font.size = Pt(font_size_pt)
    run.font.color.rgb = color
    return run


def add_logo_to_header(doc, logo_path):
    if not logo_path or not os.path.exists(logo_path):
        return
    for section in doc.sections:
        header = section.header
        for para in list(header.paragraphs[1:]):
            para._element.getparent().remove(para._element)
        first_para = header.paragraphs[0]
        first_para.clear()
        first_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = first_para.add_run()
        run.add_picture(logo_path, width=Cm(LOGO_WIDTH_CM))


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        for para in footer.paragraphs:
            para.clear()

        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        _add_purple_top_border(para)
        _set_right_tab_stop(para)
        _styled_run(para, FOOTER_TEXT)
        para.add_run("\t")
        _styled_run(para, "Pagina ")
        _add_field(para, "PAGE")
        _styled_run(para, " de ")
        _add_field(para, "NUMPAGES")


def main():
    if len(sys.argv) not in (2, 3):
        print("Uso: python3 exportar.py <arquivo.docx> [logo.png]", file=sys.stderr)
        sys.exit(1)

    docx_path = sys.argv[1]
    logo_path = sys.argv[2] if len(sys.argv) == 3 else None

    if not os.path.exists(docx_path):
        print(f"Erro: arquivo nao encontrado: {docx_path}", file=sys.stderr)
        sys.exit(1)

    doc = Document(docx_path)
    add_logo_to_header(doc, logo_path)
    add_footer(doc)
    doc.save(docx_path)
    print(f"OK: {docx_path}")


if __name__ == "__main__":
    main()
